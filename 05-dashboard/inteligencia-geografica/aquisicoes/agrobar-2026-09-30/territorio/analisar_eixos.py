#!/usr/bin/env python3
"""Primeira base regional de compras, cultura e esporte do REGIC 2018.

Não baixa arquivos, não altera a aplicação e não recalcula o índice Agrobar.
Requer openpyxl. Use: python analisar_eixos.py --cidade 4113700
As duas unidades espaciais oficiais, Município e Cidade REGIC, ficam separadas.
"""
from __future__ import annotations
import argparse, collections, csv, datetime, gzip, hashlib, io, json, math, zipfile
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent
RAW, OUT = ROOT / 'raw', ROOT / 'derived'
THEMES = {
    'Q1': {'slug': 'vestuario_calcados', 'name': 'Vestuário e calçados', 'city_ia': 'VAR57'},
    'Q2': {'slug': 'moveis_eletroeletronicos', 'name': 'Móveis e eletroeletrônicos', 'city_ia': 'VAR58'},
    'Q6': {'slug': 'cultura', 'name': 'Atividades culturais', 'city_ia': 'VAR62'},
    'Q7': {'slug': 'esporte', 'name': 'Atividades esportivas', 'city_ia': 'VAR63'},
}
FILES = [
    'REGIC2018_Municipios_Ligacoes_e_atracao_xlsx.zip',
    'REGIC2018_Municipios_Hierarquia_e_regiao.xlsx',
    'REGIC2018_Cidades_v2.xlsx', 'REGIC2018_Ligacoes_entre_Cidades.xlsx',
]

def now():
    return datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=-3))).isoformat(timespec='seconds')

def sha(path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1024**2), b''): h.update(block)
    return h.hexdigest()

def code(value):
    if value in [None, '']: return None
    if isinstance(value, (float, int)): return str(int(value))
    return str(value).strip()

def number(value):
    if value in [None, '']: return None
    result = float(value)
    assert math.isfinite(result), value
    return result

def records(source, dictionaries=None):
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        if dictionaries is not None:
            for sheet in workbook.worksheets[1:]:
                dictionaries[sheet.title] = [list(row) for row in sheet.iter_rows(values_only=True) if any(v is not None for v in row)]
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        header = next(rows)
        for row in rows:
            if row[0] is not None: yield dict(zip(header, row))
    finally:
        workbook.close()

class CsvOutput:
    def __init__(self, name, fields):
        self.path = OUT / (name + '.csv.gz')
        self.binary = self.path.open('wb')
        self.packed = gzip.GzipFile(fileobj=self.binary, mode='wb', mtime=0)
        self.stream = io.TextIOWrapper(self.packed, encoding='utf-8', newline='')
        self.writer = csv.DictWriter(self.stream, fieldnames=fields, extrasaction='raise')
        self.writer.writeheader(); self.rows = 0
    def add(self, row):
        self.writer.writerow(row); self.rows += 1
    def finish(self):
        self.stream.close(); self.binary.close()
        return dict(path=str(self.path.relative_to(ROOT)), rows=self.rows,
                    bytes=self.path.stat().st_size, sha256=sha(self.path))

def write_json(name, value):
    target = OUT / name
    target.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')
    return dict(path=str(target.relative_to(ROOT)), bytes=target.stat().st_size, sha256=sha(target))

def ranking(rows, value='ia_oficial', limit=10):
    ordered = sorted((r for r in rows if r[value] is not None and r[value] > 0), key=lambda r: (-r[value], r['codigo']))
    # Inclui empates no limite: um valor idêntico nunca recebe posição diferente.
    result = []; previous = None; rank = 0
    for position, row in enumerate(ordered, 1):
        if row[value] != previous: rank = position; previous = row[value]
        if limit is not None and rank > limit: break
        result.append(dict(posicao=rank, **row))
    return result

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cidade', action='append', default=[], help='Código IBGE municipal para relatório auditável; pode repetir.')
    args = parser.parse_args(); OUT.mkdir(exist_ok=True)
    summary = dict(generated_at=now(), reference=2018, status='em_revisao', themes=THEMES, sources=[], outputs={}, checks=[], limitations=[
        'Q1/Q2/Q6/Q7 municipais são percentuais estimados publicados, descritos como percentuais de citação na nota oficial; não são gasto, receita nem participação de mercado.',
        'quest_1/2/6/7 de Cidades são ordens 1, 2 ou 3; zero indica ausência de ligação nessa temática na linha, não ausência de consumo.',
        'Municípios e Cidades REGIC são unidades diferentes. A Cidade pode reunir municípios em um arranjo; suas relações internas não integram a rede interurbana principal.',
        'IA é o índice de atração oficial. Não é contagem observada de pessoas, viagens ou consumidores. Não foram criados pesos nem recalculadas notas.',
        'Referência 2018: não representa automaticamente comportamento atual. Cultura e esporte não medem fluxo para bares, consumo noturno ou preferência de marca.',
        'Municípios sem questionário ou sem destino citado não recebem valores inventados. Somas de percentuais fora de 100 são registradas sem normalização.',
        'Estatísticas regionais abaixo são contagens de ligações e códigos únicos por UF/Grande Região; os recortes se sobrepõem e não devem ser somados entre níveis.',
        'IA de Cidade/arranjo é mantido uma vez por Cidade, sem copiá-lo como valor municipal. Populações não são somadas nas ligações.',
        'Comparação com operações Fancore e interface ficam para uma etapa posterior.',
        'Destino 0, Outros / Não circula jornal, é categoria oficial não geográfica. É conservada nos percentuais e conciliações, mas não vira município, polo ou fluxo para fora da região.',
    ])
    for name in FILES:
        path = RAW / name; sidecar = json.loads((RAW / (name + '.download.json')).read_text())
        digest = sha(path)
        assert digest == sidecar['sha256'] and path.stat().st_size == sidecar['bytes'], name
        summary['sources'].append(dict(path=str(path.relative_to(ROOT)), url=sidecar['url'],
                                       fetched_at=sidecar['fetched_at'], bytes=path.stat().st_size, sha256=digest))
    dictionaries = {}
    municipalities, members = {}, collections.defaultdict(list)
    for row in records(RAW / FILES[1]):
        ident, city = code(row['codmun']), code(row['codcid'])
        assert ident not in municipalities and len(ident) == 7
        municipalities[ident] = dict(codigo=ident, nome=row['Município'], uf=row['siguf'], grande_regiao=row['Grande Região'],
            codigo_cidade_regic=city, tipo_cidade_regic=row['Categoria'], codigo_arranjo=code(row['codap']),
            nome_arranjo=row['Arranjo Populacional'], hierarquia_cidade=row['Hierarquia'],
            regioes_influencia=[row[f'Região de influência {i}'] for i in range(1, 5) if row[f'Região de influência {i}']])
        members[city].append(ident)
    assert len(municipalities) == 5570 and len(members) == 4899
    assert len({member for group in members.values() for member in group}) == sum(map(len, members.values())) == 5570
    selected = [code(i) for i in args.cidade]
    assert all(i in municipalities for i in selected), 'Código de município ausente do REGIC 2018'
    cities = {}
    for row in records(RAW / FILES[2], dictionaries.setdefault('cidades', {})):
        ident = code(row['COD_CIDADE']); assert ident in members and ident not in cities
        cities[ident] = dict(codigo=ident, nome=row['NOME_CIDADE'], uf_nucleo=row['UF'],
            grande_regiao_nucleo=municipalities[ident]['grande_regiao'],
            tipo=municipalities[ident]['tipo_cidade_regic'], municipios=sorted(members[ident]),
            ufs=sorted({municipalities[m]['uf'] for m in members[ident]}),
            hierarquia=row['VAR10'], nivel=row['VAR09'], ia_geral=number(row['VAR56']),
            atracao={q: number(row[t['city_ia']]) for q, t in THEMES.items()})
    assert set(cities) == set(members)
    for name, data in [('eixos-municipios', municipalities), ('eixos-cidades', cities)]:
        fields = list(next(iter(data.values())))
        output = CsvOutput(name, fields)
        for ident, row in sorted(data.items()):
            output.add({k: json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v for k, v in row.items()})
        summary['outputs'][name] = output.finish()

    municipal_edges = {q: [] for q in THEMES}; city_edges = {q: [] for q in THEMES}
    q_sums, ia_sums = collections.defaultdict(float), collections.defaultdict(float)
    incoming = collections.defaultdict(set); outgoing = collections.defaultdict(set)
    origins, destinations, pairs = set(), set(), set()
    municipal_attraction = {}; non_municipal_attraction = {}
    with zipfile.ZipFile(RAW / FILES[0]) as archive:
        assert archive.testzip() is None
        pdf_name = next(n for n in archive.namelist() if n.lower().endswith('.pdf'))
        pdf = OUT / 'eixos-fonte-metodologia-regic.pdf'; pdf.write_bytes(archive.read(pdf_name))
        summary['outputs']['metodologia_oficial'] = dict(path=str(pdf.relative_to(ROOT)), archive_member=pdf_name, bytes=pdf.stat().st_size, sha256=sha(pdf))
        attraction_file = 'REGIC2018_Quest_Atracao_Municipios.xlsx'
        for row in records(io.BytesIO(archive.read(attraction_file)), dictionaries.setdefault('atracao_municipal', {})):
            ident = code(row['CODMUN']); assert ident in municipalities or ident == '0'
            target = municipal_attraction if ident in municipalities else non_municipal_attraction
            assert ident not in target
            target[ident] = dict(nome=row['NOME_MUN'], ia_geral=number(row['IA']), temas={q: number(row['IA_' + q]) for q in THEMES})
        output = CsvOutput('eixos-ligacoes-municipais', ['origem', 'nome_origem', 'uf_origem', 'destino', 'nome_destino', 'uf_destino',
            'cidade_regic_origem', 'cidade_regic_destino', 'mesma_cidade_regic', 'tema', 'quesito', 'percentual_estimado_oficial',
            'ia_oficial_ligacao', 'destino_tipo', 'status', 'referencia'])
        for row in records(io.BytesIO(archive.read('REGIC2018_Quest_Ligacoes_entre_Municipios.xlsx')), dictionaries.setdefault('ligacoes_municipais', {})):
            origin, destination = code(row['MUN_ORIGEM']), code(row['MUN_DESTINO'])
            assert origin in municipalities and (destination in municipalities or destination in non_municipal_attraction) and (origin, destination) not in pairs
            pairs.add((origin, destination)); origins.add(origin); destinations.add(destination)
            om = municipalities[origin]
            dm = municipalities.get(destination, dict(nome=row['NOME_DESTINO'], uf=None, codigo_cidade_regic=None))
            for q, theme in THEMES.items():
                value, ia = number(row[q]), number(row['IA_' + q])
                assert value is None or 0 <= value <= 100
                assert ia is None or ia >= 0
                record = dict(origem=origin, nome_origem=om['nome'], uf_origem=om['uf'], destino=destination,
                    nome_destino=dm['nome'], uf_destino=dm['uf'], cidade_regic_origem=om['codigo_cidade_regic'],
                    cidade_regic_destino=dm['codigo_cidade_regic'], mesma_cidade_regic=(om['codigo_cidade_regic'] == dm['codigo_cidade_regic']) if dm['codigo_cidade_regic'] else None,
                    tema=theme['slug'], quesito=q, percentual_estimado_oficial=value, ia_oficial_ligacao=ia,
                    destino_tipo='municipio' if destination in municipalities else 'categoria_nao_geografica',
                    status='ausente' if value is None else 'citado' if value > 0 else 'zero_publicado', referencia=2018)
                output.add(record)
                if value is not None: q_sums[(origin, q)] += value
                if ia is not None: ia_sums[(destination, q)] += ia
                if value is not None and value > 0:
                    municipal_edges[q].append(record); incoming[(destination, q)].add(origin); outgoing[(origin, q)].add(destination)
        summary['outputs']['eixos-ligacoes-municipais'] = output.finish()
    assert len(origins) == 5503 and len(municipal_attraction) == 5276 and set(non_municipal_attraction) == {'0'}
    assert destinations == set(municipal_attraction) | set(non_municipal_attraction)
    differences = []
    for ident, data in (municipal_attraction | non_municipal_attraction).items():
        for q in THEMES:
            observed, official = ia_sums[(ident, q)], data['temas'][q]
            if official is not None and not math.isclose(observed, official, rel_tol=1e-10, abs_tol=1e-6):
                differences.append(dict(codigo=ident, quesito=q, soma_ligacoes=observed, ia_atracao_oficial=official, diferenca=observed-official))
    summary['ia_reconciliation'] = dict(compared=(len(municipal_attraction) + len(non_municipal_attraction)) * len(THEMES), discrepancies=len(differences), examples=differences[:20])
    summary['non_geographic_destinations'] = non_municipal_attraction
    # A diferença, se existir na fonte, deve ser exposta. Nunca substitui IA oficial.
    summary['outputs']['eixos-divergencias-ia'] = write_json('eixos-divergencias-ia.json', differences)
    sums_output = CsvOutput('eixos-somas-por-origem', ['origem', 'nome', 'quesito', 'soma_percentuais_oficiais', 'situacao'])
    sum_status = collections.defaultdict(collections.Counter)
    for (ident, q), value in sorted(q_sums.items()):
        status = '100_com_tolerancia_1e-6' if abs(value-100) <= 1e-6 else 'zero' if value == 0 else 'fora_de_100'
        sum_status[q][status] += 1
        sums_output.add(dict(origem=ident, nome=municipalities[ident]['nome'], quesito=q, soma_percentuais_oficiais=value, situacao=status))
    summary['outputs']['eixos-somas-por-origem'] = sums_output.finish()
    summary['percentage_sums'] = {q: dict(counts) for q, counts in sum_status.items()}
    summary['municipal_coverage'] = dict(municipios_cadastro=len(municipalities), origens_questionario=len(origins),
        destinos_municipais_citados=len(municipal_attraction), destinos_nao_geograficos=len(non_municipal_attraction), pares_todos_temas=len(pairs),
        municipios_sem_questionario=sorted(municipalities.keys() - origins),
        municipios_sem_atracao_publicada=sorted(municipalities.keys() - municipal_attraction.keys()))

    city_pairs = set(); city_incoming = collections.defaultdict(set)
    city_records = {}; duplicate_city_records = []
    output = CsvOutput('eixos-ligacoes-cidades', ['origem', 'nome_origem', 'tipo_origem', 'destino', 'nome_destino', 'tipo_destino',
        'tema', 'campo_oficial', 'ordem_oficial', 'status', 'distancia_reta_km', 'vinculo_rede_urbana', 'ids_ligacoes_fonte_json', 'referencia'])
    for row in records(RAW / FILES[3], dictionaries.setdefault('ligacoes_cidades', {})):
        origin, destination = code(row['cod_ori']), code(row['cod_dest'])
        assert origin in cities and destination in cities
        pair = (origin, destination)
        if pair in city_records:
            previous = city_records[pair]
            duplicate_city_records.append(dict(origem=origin, destino=destination, registros=[dict(previous), dict(row)]))
            assert previous['dist_km'] == row['dist_km'] and previous['vinculo'] == row['vinculo']
            for q in THEMES:
                field = 'quest_' + q[1:]
                positive = {int(v) for v in [previous[field], row[field]] if v not in [None, '', '0', 0]}
                assert len(positive) <= 1, ('Ordens positivas conflitantes na fonte', pair, field, positive)
                previous[field] = str(next(iter(positive))) if positive else previous[field]
            previous['ids_ligacoes_fonte'].append(row['id_reg'])
        else:
            city_records[pair] = dict(row, ids_ligacoes_fonte=[row['id_reg']])
    for (origin, destination), row in sorted(city_records.items()):
        city_pairs.add((origin, destination))
        assert origin != destination
        for q, theme in THEMES.items():
            field = 'quest_' + q[1:]; rank = int(row[field]) if row[field] not in [None, ''] else None
            assert rank in [None, 0, 1, 2, 3]
            record = dict(origem=origin, nome_origem=cities[origin]['nome'], tipo_origem=cities[origin]['tipo'],
                destino=destination, nome_destino=cities[destination]['nome'], tipo_destino=cities[destination]['tipo'],
                tema=theme['slug'], campo_oficial=field, ordem_oficial=rank,
                status='ausente' if rank is None else 'indicado' if rank > 0 else 'zero_publicado',
                distancia_reta_km=number(row['dist_km']), vinculo_rede_urbana=row['vinculo'],
                ids_ligacoes_fonte_json=json.dumps(row['ids_ligacoes_fonte']), referencia=2018)
            output.add(record)
            if rank is not None and rank > 0:
                city_edges[q].append(record); city_incoming[(destination, q)].add(origin)
    summary['outputs']['eixos-ligacoes-cidades'] = output.finish()
    summary['outputs']['duplicidades_fonte_cidades'] = write_json('eixos-duplicidades-fonte-cidades.json', duplicate_city_records)
    summary['city_coverage'] = dict(cidades=len(cities), municipios_isolados=sum(c['tipo'] == 'Município isolado' for c in cities.values()),
        arranjos=sum(c['tipo'] == 'Arranjo Populacional' for c in cities.values()), pares_todas_naturezas=len(city_pairs),
        ligacoes_por_tema={q: len(e) for q, e in city_edges.items()}, pares_repetidos_na_fonte=len(duplicate_city_records),
        tratamento_duplicidades='União das indicações por par/tema: zero não apaga ordem positiva; ordens positivas conflitantes interrompem execução; IDs originais preservados.')

    poles = {'municipios': {}, 'cidades': {}}
    for scale in poles:
        output = CsvOutput('eixos-polos-' + scale, ['codigo', 'nome', 'uf', 'grande_regiao', 'tipo', 'codigo_cidade_regic',
            'tema', 'campo_oficial_ia', 'ia_oficial', 'origens_distintas_com_ligacao', 'origens_externas_cidade_regic', 'cidades_regic_externas_distintas', 'referencia'])
        for q, theme in THEMES.items():
            entries = []
            for ident, value in sorted(municipal_attraction.items() if scale == 'municipios' else cities.items()):
                if scale == 'municipios':
                    loc = municipalities[ident]; ia = value['temas'][q]
                    record = dict(codigo=ident, nome=loc['nome'], uf=loc['uf'], grande_regiao=loc['grande_regiao'],
                        tipo='Município', codigo_cidade_regic=loc['codigo_cidade_regic'], tema=theme['slug'],
                        campo_oficial_ia='IA_' + q, ia_oficial=ia, origens_distintas_com_ligacao=len(incoming[(ident, q)]),
                        origens_externas_cidade_regic=sum(municipalities[i]['codigo_cidade_regic'] != loc['codigo_cidade_regic'] for i in incoming[(ident, q)]),
                        cidades_regic_externas_distintas=len({municipalities[i]['codigo_cidade_regic'] for i in incoming[(ident, q)] if municipalities[i]['codigo_cidade_regic'] != loc['codigo_cidade_regic']}), referencia=2018)
                else:
                    record = dict(codigo=ident, nome=value['nome'], uf=value['uf_nucleo'], grande_regiao=value['grande_regiao_nucleo'],
                        tipo=value['tipo'], codigo_cidade_regic=ident, tema=theme['slug'], campo_oficial_ia=theme['city_ia'],
                        ia_oficial=value['atracao'][q], origens_distintas_com_ligacao=len(city_incoming[(ident, q)]),
                        origens_externas_cidade_regic=len(city_incoming[(ident, q)]), cidades_regic_externas_distintas=len(city_incoming[(ident, q)]), referencia=2018)
                output.add(record); entries.append(record)
            poles[scale][q] = entries
        summary['outputs']['eixos-polos-' + scale] = output.finish()
    summary['top_poles'] = {scale: {q: ranking(rows) for q, rows in themes.items()} for scale, themes in poles.items()}

    regional = CsvOutput('eixos-regioes', ['nivel', 'regiao', 'quesito', 'tema', 'municipios_cadastro', 'origens_questionario',
        'origens_com_citacao', 'destinos_citados_na_regiao', 'pares_com_origem_na_regiao', 'pares_internos', 'pares_para_fora',
        'pares_destino_nao_geografico', 'pares_recebidos_de_fora', 'principais_polos_municipais_json', 'referencia'])
    for field in ['grande_regiao', 'uf']:
        for region in sorted({m[field] for m in municipalities.values()}):
            ids = {i for i, m in municipalities.items() if m[field] == region}
            for q, theme in THEMES.items():
                edges = municipal_edges[q]; starts = [e for e in edges if e['origem'] in ids]
                internal = sum(e['destino'] in ids for e in starts)
                unmapped = sum(e['destino'] not in municipalities for e in starts)
                row = dict(nivel=field, regiao=region, quesito=q, tema=theme['slug'], municipios_cadastro=len(ids),
                    origens_questionario=len(ids & origins), origens_com_citacao=len({e['origem'] for e in starts}),
                    destinos_citados_na_regiao=len({e['destino'] for e in edges if e['destino'] in ids}),
                    pares_com_origem_na_regiao=len(starts), pares_internos=internal, pares_para_fora=len(starts)-internal-unmapped,
                    pares_destino_nao_geografico=unmapped,
                    pares_recebidos_de_fora=sum(e['destino'] in ids and e['origem'] not in ids for e in edges),
                    principais_polos_municipais_json=json.dumps(ranking([p for p in poles['municipios'][q] if p['codigo'] in ids]), ensure_ascii=False), referencia=2018)
                assert row['pares_internos'] + row['pares_para_fora'] + row['pares_destino_nao_geografico'] == row['pares_com_origem_na_regiao']
                regional.add(row)
    summary['outputs']['eixos-regioes'] = regional.finish()
    for ident in selected:
        loc = municipalities[ident]; cid = loc['codigo_cidade_regic']
        report = dict(reference=2018, municipio=loc, cidade_regic=cities[cid],
            questionario_municipal_aplicado=ident in origins, atracao_municipal_publicada=ident in municipal_attraction,
            status_destinos_municipais='dados_publicados' if ident in origins else 'questionario_nao_aplicado',
            limites=summary['limitations'], temas={})
        for q, theme in THEMES.items():
            report['temas'][q] = dict(nome=theme['name'],
                municipio_ia_oficial=municipal_attraction.get(ident, {}).get('temas', {}).get(q),
                destinos_municipais=sorted([e for e in municipal_edges[q] if e['origem'] == ident], key=lambda e: (-e['percentual_estimado_oficial'], e['destino'])),
                origens_municipais=sorted([e for e in municipal_edges[q] if e['destino'] == ident], key=lambda e: (-e['ia_oficial_ligacao'], e['origem'])),
                destinos_cidade_regic=sorted([e for e in city_edges[q] if e['origem'] == cid], key=lambda e: (e['ordem_oficial'], e['destino'])),
                origens_cidade_regic=sorted([e for e in city_edges[q] if e['destino'] == cid], key=lambda e: (e['ordem_oficial'], e['origem'])))
        summary['outputs']['cidade_' + ident] = write_json('eixos-cidade-' + ident + '.json', report)
    summary['outputs']['dicionarios'] = write_json('eixos-dicionarios-oficiais.json', dictionaries)
    summary['checks'] = [
        'Quatro fontes comparadas com os SHA-256 e tamanhos dos downloads oficiais; ZIP com CRC válido.',
        '5.570 municípios únicos pertencem a exatamente uma das 4.899 Cidades REGIC; nenhuma população foi somada nos vínculos.',
        'Pares origem/destino únicos em cada escala; quatro linhas temáticas por par preservam zeros e ausências.',
        'Percentuais municipais entre 0 e 100; ordens de Cidades limitadas a 0, 1, 2, 3; IA não negativo.',
        'Somas de percentuais por origem registradas, sem normalização; soma dos IA de ligações confrontada com o IA municipal oficial.',
        'Contagens regionais de ligações internas e externas conciliadas; polos ordenados por IA oficial com empates preservados.',
    ]
    summary['municipal_coverage']['positive_links_by_theme'] = {q: len(e) for q, e in municipal_edges.items()}
    path = ROOT / 'eixos-verificacao.json'; path.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({k: summary[k] for k in ['city_coverage', 'percentage_sums', 'ia_reconciliation']}, ensure_ascii=False, indent=2))
    print('Arquivos produzidos:', len(summary['outputs']), '| relatório:', path)

if __name__ == '__main__': main()
