#!/usr/bin/env python3
"""Extrai agregados oficiais sem alterar fórmulas do Geo ou imputar ausências.

Série Caged consolidada na edição agosto/2026, com e sem ajustes separados.
RAIS 2023 a 2025 usa a revisão mais recente, publicada no XLSX de 2025.
"""
import json
import re
from pathlib import Path
import unicodedata

import openpyxl
import pandas as pd
from coletar import ROOT, now, sha

PROJECT = ROOT.parents[4]
CITIES = PROJECT / '05-dashboard/demo/site/inteligencia-geografica/dados/municipios.json'
MONTHS = ['janeiro', 'fevereiro', 'marco', 'abril', 'maio', 'junho', 'julho', 'agosto', 'setembro', 'outubro', 'novembro', 'dezembro']
FIELDS = {'Estoque': 'estoque', 'Admissões': 'admissoes', 'Desligamentos': 'desligamentos', 'Saldos': 'saldo', 'Variação Relativa (%)': 'variacao_percentual'}


def normalized(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(value or '').lower().strip()) if not unicodedata.combining(c))


def month(value):
    match = re.fullmatch(r'([a-z]+)/([0-9]{4})', normalized(value))
    return f'{match[2]}-{MONTHS.index(match[1])+1:02d}' if match and match[1] in MONTHS else None


def number(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def emit(frame, name, report, **metadata):
    path = ROOT / 'derived' / (name + '.parquet')
    frame.to_parquet(path, index=False)
    report['datasets'].append({'path': str(path.relative_to(ROOT)), 'rows': len(frame), 'columns': list(frame.columns),
                               'bytes': path.stat().st_size, 'sha256': sha(path), **metadata})
    print(name, len(frame), metadata.get('municipalities'), flush=True)


def main():
    crosswalk = {str(x['id'])[:6]: x for x in json.loads(CITIES.read_text())['cities']}
    assert len(crosswalk) == 5571
    report = {'generated_at': now(), 'datasets': [], 'checks': [], 'source_notes': [],
              'crosswalk': {'path': str(CITIES.relative_to(PROJECT)), 'sha256': sha(CITIES), 'method': 'Código MTE de seis dígitos associado ao prefixo IBGE de sete dígitos; sem gerar dígito verificador.'}}
    wb = openpyxl.load_workbook(ROOT / 'raw/rais-2025-tabelas.xlsx', read_only=True, data_only=True)
    rows = list(wb['TABELA 4'].values)
    sectors = []
    sector = None
    for column, (label, year) in enumerate(zip(rows[12], rows[13])):
        if label in ('Agropecuária', 'Indústria', 'Construção', 'Comércio', 'Serviços', 'Total'):
            sector = label
        if isinstance(year, int) and year in (2023, 2024, 2025):
            sectors.append((column, sector, year))
    result = []
    notes = []
    for index, row in enumerate(rows[15:], 16):
        if not isinstance(row[2], (int, float)):
            if any(isinstance(x, str) and x for x in row):
                notes.append({'excel_row': index, 'text': [x for x in row if isinstance(x, str)]})
            continue
        code = str(int(row[2]))
        city = crosswalk.get(code)
        for column, sector, year in sectors:
            result.append({'municipio_mte': code, 'municipio_ibge': str(city['id']) if city else None,
                           'uf': row[1], 'municipio_fonte': row[3], 'ano': year, 'setor': sector,
                           'vinculos_31_dezembro': number(row[column]), 'source_id': 'rais-2025-tabelas', 'source_row': index})
    frame = pd.DataFrame(result)
    assert not frame.duplicated(['municipio_mte', 'ano', 'setor']).any()
    total = frame[frame.setor.eq('Total')].set_index(['municipio_mte', 'ano']).vinculos_31_dezembro
    summed = frame[~frame.setor.eq('Total')].groupby(['municipio_mte', 'ano']).vinculos_31_dezembro.sum(min_count=1)
    mismatch = (summed - total).fillna(0).ne(0)
    differences = [{'municipio_mte': code, 'ano': year, 'sum_sectors': int(summed.loc[(code, year)]),
                    'source_total': int(total.loc[(code, year)]), 'difference': int(total.loc[(code, year)] - summed.loc[(code, year)])}
                   for code, year in mismatch[mismatch].index]
    report['checks'].append({'check': 'rais_setores_somam_total_municipal', 'passed': not differences,
                             'municipality_years': len(total), 'differences': differences,
                             'interpretation': 'Divergências já presentes na tabela oficial; total e setores preservados sem redistribuir vínculos.'})
    emit(frame, 'rais-vinculos-municipio-setor-2023-2025', report,
         municipalities=frame.municipio_ibge.nunique(), geographic_records=frame.municipio_mte.nunique(), years=sorted(frame.ano.unique().tolist()),
         unmatched_codes=sorted(frame.loc[frame.municipio_ibge.isna(), 'municipio_mte'].unique().tolist()),
         units={'vinculos_31_dezembro': 'vínculos formais ativos em 31/12; não pessoas distintas'},
         vintage='RAIS 2025', source_id='rais-2025-tabelas')
    report['source_notes'].append({'source_id': 'rais-2025-tabelas', 'sheet': 'TABELA 4', 'notes': notes})
    # Salários agregados no XLSX são contexto nacional/UF/setorial, não municipais.
    salary_rows = list(wb['TABELA 6'].values)
    salary_data = []
    category = 'Brasil'
    for index, row in enumerate(salary_rows, 1):
        if index < 13:
            continue
        if row[1] and row[2] is None and not isinstance(row[3], (int, float)):
            category = str(row[1])
        if not isinstance(row[3], (int, float)):
            continue
        label = row[2] or 'Brasil'
        for column, year in [(3, 2023), (4, 2024), (5, 2025)]:
            salary_data.append({'recorte': category, 'categoria': label, 'ano': year,
                                'remuneracao_dezembro_reais': number(row[column]), 'source_row': index})
    emit(pd.DataFrame(salary_data), 'rais-remuneracao-contexto-2023-2025', report, source_id='rais-2025-tabelas',
         geographic_scope='Brasil, regiões e UFs; inclui categorias setoriais nacionais. NÃO é remuneração municipal.',
         limitations=['O sumário chama de nominal, mas o corpo da tabela chama de Real. Preservado o valor sem redeflacionar; confirmar nota técnica e base de preços antes de cruzar.'])
    wb.close()
    wb = openpyxl.load_workbook(ROOT / 'raw/caged-202608.xlsx', read_only=True, data_only=True)
    for sheet_name, adjusted in [('Tabela 8.1', True), ('Tabela 8', False)]:
        sheet = wb[sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        heads = [next(iterator) for _ in range(6)]
        cols = []
        current = None
        for col, (label, metric) in enumerate(zip(heads[4], heads[5])):
            if label is not None:
                current = month(label)
            if current and metric in FIELDS:
                cols.append((col, current, FIELDS[metric]))
        results = []
        notes = []
        for index, row in enumerate(iterator, 7):
            if not isinstance(row[2], (int, float)):
                if any(isinstance(x, str) and x for x in row[:5]):
                    notes.append({'excel_row': index, 'text': [x for x in row[:5] if isinstance(x, str)]})
                continue
            code = str(int(row[2]))
            city = crosswalk.get(code)
            periods = {}
            for col, period, metric in cols:
                if period not in periods:
                    periods[period] = {'municipio_mte': code, 'municipio_ibge': str(city['id']) if city else None,
                                       'uf': row[1], 'municipio_fonte': row[3], 'competencia': period,
                                       'com_ajustes': adjusted, 'edicao': '2026-08', 'source_row': index}
                periods[period][metric] = number(row[col])
            results.extend(periods.values())
        df = pd.DataFrame(results)
        assert not df.duplicated(['municipio_mte', 'competencia']).any()
        diff = df.admissoes - df.desligamentos - df.saldo
        bad = df.loc[diff.fillna(0).ne(0)]
        assert bad.empty, f'{sheet_name}: admissões menos desligamentos difere do saldo'
        report['checks'].append({'check': 'caged_admissoes_menos_desligamentos_igual_saldo', 'sheet': sheet_name,
                                 'passed': True, 'rows': len(df)})
        national_sheet = wb['Tabela 5.1' if adjusted else 'Tabela 5']
        official = {}
        for row in national_sheet.iter_rows(min_row=6, max_col=7, values_only=True):
            period = month(row[1])
            if period:
                official[period] = dict(zip(['estoque', 'admissoes', 'desligamentos', 'saldo'], row[2:6]))
        national_diff = []
        for period, group in df.groupby('competencia'):
            if period in official:
                for key in ['estoque', 'admissoes', 'desligamentos', 'saldo']:
                    summed = int(group[key].sum())
                    if summed != official[period][key]:
                        national_diff.append({'competencia': period, 'metric': key, 'sum_municipal_rows': summed,
                                              'official_brasil': official[period][key], 'difference': summed - official[period][key]})
        # Diferenças são preservadas para auditoria, nunca corrigidas ou imputadas.
        report['checks'].append({'check': 'caged_municipios_contra_total_brasil', 'sheet': sheet_name,
                                 'passed': not national_diff, 'differences': national_diff,
                                 'interpretation': 'Diferenças históricas somente em estoque; a nota oficial informa que estoques municipais negativos são mostrados como zero, mas considerados nos totais. Valores originais preservados.'})
        latest = df[df.competencia.eq('2026-08')]
        observed = {key: int(latest[key].sum()) for key in ['estoque', 'admissoes', 'desligamentos', 'saldo']}
        assert observed == official['2026-08'], 'Agosto/2026 não conciliou com Brasil'
        report['checks'].append({'check': 'caged_2026_08_conciliado_brasil', 'sheet': sheet_name, 'passed': True, 'totals': observed})
        suffix = 'com-ajustes' if adjusted else 'sem-ajustes'
        emit(df, 'caged-municipio-2020-2026-' + suffix, report, source_id='caged-202608', sheet=sheet_name,
             municipalities=df.municipio_ibge.nunique(), geographic_records=df.municipio_mte.nunique(), months=df.competencia.nunique(), start=df.competencia.min(), end=df.competencia.max(),
             unmatched_codes=sorted(df.loc[df.municipio_ibge.isna(), 'municipio_mte'].unique().tolist()),
             missing_numeric={key: int(df[key].isna().sum()) for key in ['estoque', 'admissoes', 'desligamentos', 'saldo']},
             units={'estoque': 'vínculos', 'admissoes': 'movimentações', 'desligamentos': 'movimentações', 'saldo': 'vínculos', 'variacao_percentual': '%'})
        last24 = df[df.competencia.ge('2024-09')].copy()
        csv = ROOT / 'derived' / ('caged-municipio-24meses-' + suffix + '.csv.gz')
        last24.to_csv(csv, index=False, compression={'method': 'gzip', 'mtime': 0})
        report['datasets'].append({'path': str(csv.relative_to(ROOT)), 'rows': len(last24), 'bytes': csv.stat().st_size,
                                   'sha256': sha(csv), 'start': '2024-09', 'end': '2026-08', 'months': last24.competencia.nunique(),
                                   'source_id': 'caged-202608', 'sheet': sheet_name})
        report['source_notes'].append({'source_id': 'caged-202608', 'sheet': sheet_name, 'notes': notes})
    wages = []
    wage_notes = []
    for index, row in enumerate(wb['Tabela 9'].iter_rows(min_row=6, max_col=5, values_only=True), 6):
        period = month(row[1])
        if period:
            wages.append({'competencia': period, 'salario_medio_real_admissao': number(row[2]),
                          'salario_medio_real_desligamento': number(row[3]), 'source_row': index})
        elif row[1]:
            wage_notes.append(str(row[1]))
    emit(pd.DataFrame(wages), 'caged-salarios-brasil-2020-2026', report, source_id='caged-202608', sheet='Tabela 9',
         geographic_scope='Brasil, sem detalhamento municipal', units='R$ reais segundo nota da fonte', notes=wage_notes)
    wb.close()
    out = ROOT / 'derived/validacao-agregados.json'
    out.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=int) + '\n')
    print('report', out, flush=True)


if __name__ == '__main__':
    main()
